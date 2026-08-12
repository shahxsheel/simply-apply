"""ATS-safe DOCX renderer using a Jake's Resume-inspired one-page layout.

The visual system follows Jake Gutierrez's MIT-licensed LaTeX template: centered contact
header, compact serif typography, ruled uppercase sections, dense bullets, and paired
role/date plus employer/location lines. It stays single-column and table-free so Word
ATS parsers read the XML in the intended order.

Source inspiration: https://github.com/jakegut/resume (MIT License).
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.schemas import StructuredResume

BODY_FONT = "Times New Roman"
BODY_SIZE = Pt(9.5)
NAME_SIZE = Pt(23)
HEADING_SIZE = Pt(11.5)
INK = RGBColor(0, 0, 0)
PAGE_WIDTH = Inches(8.5)
PAGE_HEIGHT = Inches(11)
MARGIN = Inches(0.5)
CONTENT_WIDTH = Inches(7.5)


def _set_font(run, *, size=None, bold=None, italic=None) -> None:
    run.font.name = BODY_FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), BODY_FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), BODY_FONT)
    run.font.color.rgb = INK
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def _configure_document(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    normal.font.color.rgb = INK
    normal.element.get_or_add_rPr().rFonts.set(qn("w:ascii"), BODY_FONT)
    normal.element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), BODY_FONT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    bullet = document.styles["List Bullet"]
    bullet.font.name = BODY_FONT
    bullet.font.size = BODY_SIZE
    bullet.element.get_or_add_rPr().rFonts.set(qn("w:ascii"), BODY_FONT)
    bullet.element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), BODY_FONT)
    bullet.paragraph_format.left_indent = Inches(0.22)
    bullet.paragraph_format.first_line_indent = Inches(-0.14)
    bullet.paragraph_format.space_after = Pt(0)
    bullet.paragraph_format.line_spacing = 1.0

    if "Resume Section" not in document.styles:
        section_style = document.styles.add_style(
            "Resume Section", WD_STYLE_TYPE.PARAGRAPH
        )
    else:
        section_style = document.styles["Resume Section"]
    section_style.font.name = BODY_FONT
    section_style.font.size = HEADING_SIZE
    section_style.font.bold = True
    section_style.font.color.rgb = INK
    section_style.element.get_or_add_rPr().rFonts.set(qn("w:ascii"), BODY_FONT)
    section_style.element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), BODY_FONT)
    section_style.paragraph_format.space_before = Pt(5)
    section_style.paragraph_format.space_after = Pt(2)
    section_style.paragraph_format.keep_with_next = True

    for section in document.sections:
        section.start_type = WD_SECTION.CONTINUOUS
        section.page_width = PAGE_WIDTH
        section.page_height = PAGE_HEIGHT
        section.top_margin = MARGIN
        section.bottom_margin = MARGIN
        section.left_margin = MARGIN
        section.right_margin = MARGIN
        section.header_distance = Inches(0.2)
        section.footer_distance = Inches(0.2)
        # Word's "Justified" vertical page alignment distributes the document's
        # existing paragraph spacing between equal top and bottom margins. This keeps a
        # short resume from looking top-heavy without changing ATS reading order.
        properties = section._sectPr
        existing = properties.find(qn("w:vAlign"))
        if existing is not None:
            properties.remove(existing)
        vertical_alignment = OxmlElement("w:vAlign")
        vertical_alignment.set(qn("w:val"), "both")
        properties.append(vertical_alignment)


def _bottom_rule(paragraph) -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    borders.append(bottom)


def _section_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Resume Section")
    run = paragraph.add_run(text.upper())
    _set_font(run, size=HEADING_SIZE, bold=True)
    run.font.small_caps = True
    _bottom_rule(paragraph)


def _paired_line(
    document: Document,
    left: str,
    right: str = "",
    *,
    bold: bool = False,
    italic: bool = False,
    before: float = 0,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        CONTENT_WIDTH, WD_TAB_ALIGNMENT.RIGHT
    )
    left_run = paragraph.add_run(left)
    _set_font(left_run, bold=bold, italic=italic)
    if right:
        paragraph.add_run("\t")
        right_run = paragraph.add_run(right)
        _set_font(right_run, italic=italic)


def _body_paragraph(document: Document, text: str, *, italic: bool = False) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(1)
    run = paragraph.add_run(text)
    _set_font(run, italic=italic)


def _bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.keep_together = True
    for run in paragraph.runs:
        _set_font(run)
    if not paragraph.runs:
        _set_font(paragraph.add_run(text))


def _date_range(start: str, end: str) -> str:
    if start and end:
        return f"{start} - {end}"
    return start or end or ""


def _contact_line(resume: StructuredResume) -> str:
    basics = resume.basics
    location = ", ".join(
        p for p in (basics.location.city, basics.location.region) if p
    )
    pieces = [basics.phone, basics.email, location, basics.url]
    pieces += [profile.url for profile in basics.profiles if profile.url]
    return " | ".join(piece for piece in pieces if piece)


def render_docx(resume: StructuredResume, output_path: Path) -> Path:
    document = Document()
    _configure_document(document)

    name = document.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(0)
    name_run = name.add_run(resume.basics.name or "")
    _set_font(name_run, size=NAME_SIZE, bold=True)
    name_run.font.small_caps = True

    contact = _contact_line(resume)
    if contact:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(3)
        _set_font(paragraph.add_run(contact), size=Pt(9))

    if resume.basics.summary:
        _section_heading(document, "Professional Summary")
        _body_paragraph(document, resume.basics.summary)

    if resume.work:
        _section_heading(document, "Experience")
        for job in resume.work:
            _paired_line(
                document,
                job.position or job.name,
                _date_range(job.startDate, job.endDate),
                bold=True,
                before=2,
            )
            employer = job.name if job.position else ""
            if employer or job.location:
                _paired_line(document, employer, job.location, italic=True)
            if job.summary:
                _body_paragraph(document, job.summary)
            for highlight in job.highlights:
                _bullet(document, highlight)

    if resume.education:
        _section_heading(document, "Education")
        for education in resume.education:
            _paired_line(
                document,
                education.institution,
                _date_range(education.startDate, education.endDate),
                bold=True,
                before=2,
            )
            degree = " in ".join(
                part for part in (education.studyType, education.area) if part
            )
            detail = " | ".join(
                part
                for part in (degree, f"GPA: {education.score}" if education.score else "")
                if part
            )
            if detail:
                _body_paragraph(document, detail, italic=True)
            if education.courses:
                _body_paragraph(
                    document, f"Relevant Coursework: {', '.join(education.courses)}"
                )

    if resume.projects:
        _section_heading(document, "Projects")
        for project in resume.projects:
            technologies = ", ".join(project.keywords)
            left = project.name
            if technologies:
                left = f"{left} | {technologies}" if left else technologies
            _paired_line(
                document,
                left,
                _date_range(project.startDate, project.endDate),
                bold=bool(project.name),
                before=2,
            )
            if project.description:
                _body_paragraph(document, project.description)
            for highlight in project.highlights:
                _bullet(document, highlight)

    if resume.skills:
        _section_heading(document, "Technical Skills")
        for skill in resume.skills:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(0)
            if skill.name and skill.keywords:
                _set_font(paragraph.add_run(f"{skill.name}: "), bold=True)
                _set_font(paragraph.add_run(", ".join(skill.keywords)))
            elif skill.name:
                _set_font(paragraph.add_run(skill.name))
            elif skill.keywords:
                _set_font(paragraph.add_run(", ".join(skill.keywords)))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    return output_path
